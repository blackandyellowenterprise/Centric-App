"""Data loader for educational standards from Excel files."""

import openpyxl
from typing import Dict, List
from pathlib import Path
from src.models.standards import Subject, Strand, StrandPart, GradeStandards


class StandardsLoader:
    """Loads and processes educational standards from Excel files."""

    def __init__(self, data_dir: str = "extracted_data"):
        """Initialize the standards loader.

        Args:
            data_dir: Directory containing extracted Excel files
        """
        self.data_dir = Path(data_dir)
        self.standards_cache: Dict[str, List[GradeStandards]] = {}

    def load_subject_standards(self, subject: Subject) -> List[GradeStandards]:
        """Load all standards for a given subject.

        Args:
            subject: Subject to load standards for

        Returns:
            List of GradeStandards for all grades in the subject
        """
        cache_key = subject.value
        if cache_key in self.standards_cache:
            return self.standards_cache[cache_key]

        file_mapping = {
            Subject.ELA: "Centric ELA Strands-2.xlsx",
            Subject.MATH: "Centric Math Strands-2.xlsx",
            Subject.SCIENCE: "Centric Science Strands-2.xlsx",
            Subject.SOCIAL_STUDIES: "Centric Social Studies Strands-3.xlsx"
        }

        file_path = self.data_dir / file_mapping[subject]
        if not file_path.exists():
            raise FileNotFoundError(f"Standards file not found: {file_path}")

        wb = openpyxl.load_workbook(file_path)
        all_grades = []

        for sheet_name in wb.sheetnames:
            grade_standards = self._parse_sheet(wb[sheet_name], subject)
            if grade_standards:
                all_grades.append(grade_standards)

        self.standards_cache[cache_key] = all_grades
        return all_grades

    def _parse_sheet(self, worksheet, subject: Subject) -> GradeStandards:
        """Parse a single worksheet into GradeStandards.

        Args:
            worksheet: Excel worksheet to parse
            subject: Subject of the standards

        Returns:
            GradeStandards object
        """
        # Extract grade from sheet name (e.g., "ELA 06" -> 6)
        sheet_name = worksheet.title
        try:
            grade = int(sheet_name.split()[-1])
        except (ValueError, IndexError):
            return None

        strands: Dict[str, Strand] = {}
        current_strand_title = None

        # Skip header row
        rows = list(worksheet.iter_rows(values_only=True))
        for row in rows[1:]:
            if not row or not row[1]:  # Skip empty rows
                continue

            strand_title = row[0]
            strand_code = row[1]
            prof_1 = row[2] or ""
            prof_2 = row[3] or ""
            prof_3 = row[4] or ""
            cc_alignment = row[5] or ""

            # Update current strand title
            if strand_title:
                current_strand_title = strand_title
                if current_strand_title not in strands:
                    strands[current_strand_title] = Strand(
                        title=current_strand_title,
                        parts=[]
                    )

            # Add strand part
            if current_strand_title and strand_code:
                strand_part = StrandPart(
                    code=strand_code,
                    proficiency_level_1=prof_1,
                    proficiency_level_2=prof_2,
                    proficiency_level_3=prof_3,
                    common_core_alignment=cc_alignment
                )
                strands[current_strand_title].parts.append(strand_part)

        return GradeStandards(
            subject=subject,
            grade=grade,
            strands=list(strands.values())
        )

    def get_strand_by_code(self, code: str) -> tuple[StrandPart, GradeStandards]:
        """Find a specific strand part by its code.

        Args:
            code: Strand part code (e.g., "ELA06.RLa")

        Returns:
            Tuple of (StrandPart, GradeStandards) or (None, None) if not found
        """
        # Extract subject from code
        subject_prefix = code[:3].lower()
        subject_map = {
            "ela": Subject.ELA,
            "mat": Subject.MATH,
            "sci": Subject.SCIENCE,
            "soc": Subject.SOCIAL_STUDIES
        }

        subject = subject_map.get(subject_prefix)
        if not subject:
            return None, None

        all_standards = self.load_subject_standards(subject)

        for grade_standards in all_standards:
            for strand in grade_standards.strands:
                for part in strand.parts:
                    if part.code == code:
                        return part, grade_standards

        return None, None

    def search_standards(self, query: str, subject: Subject = None, grade: int = None) -> List[tuple[StrandPart, str, Subject, int]]:
        """Search standards by keyword.

        Args:
            query: Search query
            subject: Optional subject filter
            grade: Optional grade filter

        Returns:
            List of tuples (StrandPart, strand_title, subject, grade)
        """
        results = []
        query_lower = query.lower()

        subjects_to_search = [subject] if subject else list(Subject)

        for subj in subjects_to_search:
            all_standards = self.load_subject_standards(subj)

            for grade_standards in all_standards:
                if grade and grade_standards.grade != grade:
                    continue

                for strand in grade_standards.strands:
                    for part in strand.parts:
                        # Search in all text fields
                        searchable_text = " ".join([
                            strand.title,
                            part.code,
                            part.proficiency_level_1,
                            part.proficiency_level_2,
                            part.proficiency_level_3,
                            part.common_core_alignment
                        ]).lower()

                        if query_lower in searchable_text:
                            results.append((part, strand.title, subj, grade_standards.grade))

        return results
